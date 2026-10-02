# Ghi chú nhóm: Cập nhật mô tả API và thông báo phản hồi bằng tiếng Việt.
from io import BytesIO
from pathlib import Path
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, Form, HTTPException, UploadFile, File, status
from sqlalchemy.orm import Session
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from app.models.call_record import CallRecord

from app.db.database import get_db
from app.core.config import settings
from app.services.mediafile_service import MediaFileService
from app.services.call_service import CallService
from app.schemas.mediafile_schema import SpeakerRoleUpdate
from app.api.v1.endpoints.auth import get_current_user
from app.models.user import User

router = APIRouter()


@router.get("/")
def get_media_files(
    start: Optional[str] = None,
    end: Optional[str] = None,
    offset: int = 0,
    limit: int = 20,
    operatorId: Optional[int] = None,
    searchPhrase: Optional[str] = None,
    db: Session = Depends(get_db),
):
    return MediaFileService.get_list(
        db,
        start=start,
        end=end,
        offset=offset,
        limit=limit,
        operator_id=operatorId,
        search_phrase=searchPhrase,
    )


@router.post("/upload", status_code=status.HTTP_201_CREATED)
async def upload_media_file(
    file: UploadFile = File(...),
    operatorId: Optional[int] = Form(None),
    projectId: Optional[int] = Form(None),
    clientNumber: Optional[str] = Form(None),
    createDate: Optional[datetime] = Form(None),
    db: Session = Depends(get_db),
):
    return await MediaFileService.create_from_upload(
        db,
        file=file,
        operator_id=operatorId,
        project_id=projectId,
        client_number=clientNumber,
        call_date=createDate,
    )


@router.get("/export/excel")
def export_media_files_excel(db: Session = Depends(get_db)):
    records = db.query(CallRecord).order_by(CallRecord.id.desc()).all()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Cuoc goi"

    headers = [
        "ID",
        "Tệp ghi âm",
        "Điều hành viên",
        "Dự án",
        "Số khách hàng",
        "Thời lượng (giây)",
        "Điểm tuân thủ",
        "Cảm xúc",
        "Ngày cuộc gọi",
        "Ngày tạo",
    ]
    worksheet.append(headers)

    header_fill = PatternFill("solid", fgColor="0068AD")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for record in records:
        worksheet.append(
            [
                record.id,
                Path(record.file_path).name,
                record.operator.name if record.operator else "",
                record.project.name if record.project else "",
                record.client_number or "",
                record.audio_duration,
                record.compliance_score if record.compliance_score is not None else "Chưa chấm điểm",
                record.sentiment or "",
                record.call_date,
                record.created_at,
            ]
        )

    widths = [10, 32, 24, 24, 20, 18, 18, 16, 22, 22]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="call-records.xlsx"'},
    )


@router.get("/{file_id}")
def get_media_file(file_id: int, db: Session = Depends(get_db)):
    return MediaFileService.get_by_id(db, file_id)


@router.get("/{file_id}/result")
def get_media_file_result(file_id: int, db: Session = Depends(get_db)):
    return MediaFileService.get_result(db, file_id)


@router.put("/{file_id}/speaker-roles")
def confirm_speaker_roles(
    file_id: int,
    payload: SpeakerRoleUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    del current_user
    record = CallService.confirm_speaker_roles(
        db,
        call_id=file_id,
        agent_speaker_id=payload.agentSpeakerId,
    )
    return MediaFileService.get_result(db, record.id)


@router.get("/{file_id}/stream")
def stream_media_file(file_id: int, db: Session = Depends(get_db)):
    record = db.query(CallRecord).filter(CallRecord.id == file_id).first()
    if not record:
        raise HTTPException(status_code=404, detail=f"Không tìm thấy bản ghi có mã: {file_id}")
    file_path = Path(record.file_path)
    if not file_path.is_absolute():
        file_path = settings.UPLOAD_DIR / file_path
    file_path = file_path.resolve()
    if not file_path.is_file():
        raise HTTPException(status_code=404, detail=f"Tệp âm thanh không tồn tại: {file_path}")
    return FileResponse(str(file_path), filename=file_path.name)
