import sys
import os
from io import BytesIO
from pathlib import Path
import time
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

# Add the project root directory to sys.path.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.main import app
from app.asr_client import AsrApiError
from app.db.database import SessionLocal
from app.db.init_db import initialize_database
from app.models.asr_job import AsrJob
from app.models.call_record import CallRecord
from app.models.user import User
from app.models.vocabulary import Vocabulary
from app.services.analytics_service import AnalyticsService
from app.core.security import create_access_token, decode_access_token, hash_password
from datetime import timedelta


client = TestClient(app)


def _get_auth_headers(role: str = "admin", user_id: int = 999, email: str = "admin@test.com") -> dict:
    with SessionLocal() as db:
        user = db.query(User).filter(User.id == user_id).first()
        if not user:
            user = User(
                id=user_id,
                email=email,
                password_hash=hash_password("password123"),
                full_name="Test User",
                role=role,
                is_active=True,
            )
            db.add(user)
            db.commit()
        else:
            user.role = role
            user.is_active = True
            db.commit()
    token = create_access_token({"sub": str(user_id), "email": email, "role": role})
    return {"Authorization": f"Bearer {token}"}


def test_health_check():
    res = client.get("/")
    assert res.status_code == 200
    assert res.json()["status"] == "online"


def test_legacy_api_v1_is_removed():
    assert client.get("/api/v1/auth/signin").status_code == 404
    paths = client.get("/openapi.json").json()["paths"]
    assert "/api/auth/signin" in paths
    assert not any(path.startswith("/api/v1/") for path in paths)


def test_database_initialization_is_idempotent_and_preserves_edited_regex():
    initialize_database()
    with SessionLocal() as db:
        entry = db.query(Vocabulary).filter(Vocabulary.name == "Regex | Lời chào bắt buộc").one()
        original_data = entry.data
        entry.data = {"phrases": ["custom test value"]}
        db.commit()

    initialize_database()
    with SessionLocal() as db:
        entry = db.query(Vocabulary).filter(Vocabulary.name == "Regex | Lời chào bắt buộc").one()
        assert entry.data == {"phrases": ["custom test value"]}
        entry.data = original_data
        db.commit()


def test_failed_asr_upload_is_saved_without_a_score():
    headers = _get_auth_headers(role="admin", user_id=901, email="admin_asr@test.com")
    files = {"file": ("offline-call.wav", b"audio", "audio/wav")}
    with patch(
        "app.services.asr_service.submit_to_asr",
        side_effect=AsrApiError(503, "unavailable", "ASR unavailable"),
    ):
        response = client.post("/api/transcribe/upload", files=files, headers=headers)

    assert response.status_code == 202
    assert response.json()["status"] == "failed"
    record_id = response.json()["call_record_id"]
    with SessionLocal() as db:
        record = db.query(CallRecord).filter(CallRecord.id == record_id).one()
        assert record.compliance_score is None
        dashboard = AnalyticsService.get_dashboard(db)
        assert dashboard["summaryData"]["averageNegativeLevelOverall"] is None

    result = client.get(f"/api/mediafile/{record_id}/result", headers=headers)
    assert result.status_code == 200
    assert result.json()["stt"]["text"] is None

    exported = client.get("/api/mediafile/export/excel", headers=headers)
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content), read_only=True).active
    assert any(row[0] == record_id and row[6] == "Chưa chấm điểm" for row in sheet.iter_rows(min_row=2, values_only=True))


def test_user_crud():
    admin_headers = _get_auth_headers(role="admin", user_id=902, email="admin_crud@test.com")
    username = f"telesale_{int(time.time())}"

    # 1. Create (Admin)
    user_payload = {
        "email": f"{username}@example.com",
        "password": "password123",
        "full_name": "Nguyen Van B",
        "role": "telesales",
    }

    res_create = client.post("/api/users/", json=user_payload, headers=admin_headers)
    assert res_create.status_code == 201
    user_id = res_create.json()["id"]

    # 2. Get by ID (Admin)
    res_get = client.get(f"/api/users/{user_id}", headers=admin_headers)
    assert res_get.status_code == 200
    assert res_get.json()["email"] == f"{username}@example.com"

    # 3. Update (Admin)
    res_update = client.put(
        f"/api/users/{user_id}",
        json={"full_name": "Nguyen Van B - VIP"},
        headers=admin_headers,
    )
    assert res_update.status_code == 200
    assert res_update.json()["full_name"] == "Nguyen Van B - VIP"

    # 4. List (Admin)
    res_list = client.get("/api/users/", headers=admin_headers)
    assert res_list.status_code == 200
    assert any(user["id"] == user_id for user in res_list.json())

    # 5. Telesales should NOT be able to list all users (403 Forbidden)
    staff_headers = _get_auth_headers(role="telesales", user_id=user_id, email=f"{username}@example.com")
    res_staff_list = client.get("/api/users/", headers=staff_headers)
    assert res_staff_list.status_code == 403

    # 6. Delete (Admin)
    res_delete = client.delete(f"/api/users/{user_id}", headers=admin_headers)
    assert res_delete.status_code == 200


def test_email_signin_and_expired_token():
    admin_headers = _get_auth_headers(role="admin", user_id=903, email="admin_signin@test.com")
    email = f"signin_{int(time.time())}@example.com"
    created = client.post(
        "/api/users/",
        json={
            "email": email,
            "password": "password123",
            "full_name": "Sign In Test",
            "role": "telesales",
        },
        headers=admin_headers,
    )
    assert created.status_code == 201

    response = client.post(
        "/api/auth/signin",
        json={"email": email, "password": "password123"},
    )
    assert response.status_code == 200
    assert decode_access_token(response.json()["accessToken"])["email"] == email
    assert response.json()["user"]["role"] == "telesales"

    expired = create_access_token({"sub": str(created.json()["id"])}, timedelta(seconds=-1))
    assert decode_access_token(expired) is None


def test_audio_file_upload_and_analysis_flow():
    admin_headers = _get_auth_headers(role="admin", user_id=904, email="admin_audio@test.com")

    # 1. Upload audio
    dummy_audio = b"RIFF....WAVEfmt ....data...."
    files = {
        "file": (
            "demo_call.wav",
            dummy_audio,
            "audio/wav",
        )
    }

    res_file = client.post("/api/files/upload", files=files, headers=admin_headers)
    assert res_file.status_code == 201
    file_path = res_file.json()["file_path"]

    # 2. Process AI JSON result and create a call record.
    ai_call_payload = {
        "file_path": file_path,
        "ai_result": {
            "transcript": (
                "Dạ em chào quý khách, em gọi tư vấn gói bảo hiểm. "
                "Mày không mua thì lượn đi."
            ),
            "segments": [
                {
                    "speaker": "agent",
                    "start_time": 0.0,
                    "end_time": 3.2,
                    "text": (
                        "Dạ em chào quý khách, "
                        "em gọi tư vấn gói bảo hiểm."
                    ),
                },
                {
                    "speaker": "customer",
                    "start_time": 3.5,
                    "end_time": 5.0,
                    "text": "Tôi bận lắm.",
                },
                {
                    "speaker": "agent",
                    "start_time": 5.2,
                    "end_time": 7.8,
                    "text": "Mày không mua thì lượn đi.",
                },
            ],
            "sentiment": "negative",
            "duration": 8,
        },
    }

    res_call = client.post("/api/calls/", json=ai_call_payload, headers=admin_headers)
    assert res_call.status_code == 201

    call_json = res_call.json()
    call_id = call_json["id"]

    assert call_json["compliance_score"] < 100.0
    assert len(call_json["violations"]) > 0

    # 3. Query violations
    res_violations = client.get(
        f"/api/violations/?call_record_id={call_id}",
        headers=admin_headers,
    )
    assert res_violations.status_code == 200
    assert len(res_violations.json()) > 0

    # 4. Cleanup
    res_delete = client.delete(f"/api/calls/{call_id}", headers=admin_headers)
    assert res_delete.status_code == 200
